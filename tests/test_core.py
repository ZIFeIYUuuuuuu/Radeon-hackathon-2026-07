from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
import json

from core import (
    BM25Index,
    FileMatch,
    LocalEmbeddingClient,
    LocalRetriever,
    WorkspaceScanPolicy,
    assess_file_matches,
    _direct_filename_phrase_bonus,
    _document_sections,
    compile_intent,
    discover_workspace,
    infer_intent,
    is_file_inventory_request,
    is_file_existence_request,
    markdown_brief,
    redact_sensitive_text,
    route_workspace_request,
    run_court,
    scan_sensitive_paths,
    scan_workspace,
    summarize_artifact,
)
from scripts.private_retrieval_check import sanitized_report


class ClaimCourtTests(unittest.TestCase):
    def setUp(self):
        self.retriever = LocalRetriever()
        self.retriever.index_paths((Path(__file__).parent.parent / "demo_corpus").glob("*"))

    def test_demo_case_returns_cited_verdict(self):
        evidence = self.retriever.search("Did the vendor contractually commit to 99.9% uptime?")
        _, _, verdict, _, _ = run_court("Did the vendor contractually commit to 99.9% uptime?", evidence, allow_fallback=True, demo_mode=True)
        self.assertEqual("insufficient_evidence", verdict["verdict"])
        self.assertTrue(verdict["evidence_citations"])

    def test_short_demo_corpus_does_not_repeat_terminal_chunks(self):
        self.assertGreaterEqual(len(self.retriever.evidence), 5)
        self.assertTrue(all(len(item.text) <= 480 for item in self.retriever.evidence))

    def test_unreadable_files_do_not_abort_directory_indexing(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            readable = root / "readable.md"
            unreadable = root / "empty.docx"
            readable.write_text("a readable local report", encoding="utf-8")
            unreadable.write_bytes(b"")
            retriever = LocalRetriever()
            retriever.index_paths([readable, unreadable])
        self.assertEqual(["readable.md"], [item.source for item in retriever.evidence])
        self.assertEqual(1, len(retriever.ingestion_errors))
        self.assertEqual("unreadable", retriever.ingestion_errors[0]["status"])

    def test_workspace_scanner_is_deterministic_and_explains_exclusions(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            docs = root / "docs"
            dependencies = root / "node_modules" / "package"
            hidden = root / ".private"
            docs.mkdir()
            dependencies.mkdir(parents=True)
            hidden.mkdir()
            (docs / "z-report.md").write_text("z", encoding="utf-8")
            (docs / "A-notes.txt").write_text("a", encoding="utf-8")
            (dependencies / "noise.md").write_text("not user evidence", encoding="utf-8")
            (hidden / "hidden.md").write_text("hidden", encoding="utf-8")
            (root / "binary.exe").write_bytes(b"binary")

            scan = scan_workspace(root)
            legacy_files = discover_workspace(root)

        self.assertEqual(["A-notes.txt", "z-report.md"], [path.name for path in scan.files])
        self.assertEqual(1, scan.skip_counts["excluded_directory"])
        self.assertEqual(1, scan.skip_counts["hidden_directory"])
        self.assertEqual(1, scan.skip_counts["unsupported_type"])
        self.assertEqual(list(scan.files), legacy_files)

    def test_workspace_scanner_enforces_size_and_file_count_limits(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "01-small.md").write_text("small", encoding="utf-8")
            (root / "02-too-large.md").write_text("x" * 20, encoding="utf-8")
            (root / "03-limited.md").write_text("small", encoding="utf-8")
            policy = WorkspaceScanPolicy(max_file_bytes=10, max_files=1)

            scan = scan_workspace(root, policy)

        self.assertEqual(["01-small.md"], [path.name for path in scan.files])
        self.assertEqual(1, scan.skip_counts["file_too_large"])
        self.assertEqual(1, scan.skip_counts["file_limit_reached"])
        self.assertTrue(scan.file_limit_reached)

    def test_workspace_scanner_does_not_follow_file_symlinks(self):
        with TemporaryDirectory() as directory:
            base = Path(directory)
            root = base / "workspace"
            root.mkdir()
            target = base / "outside.md"
            target.write_text("outside selected root", encoding="utf-8")
            link = root / "linked.md"
            try:
                link.symlink_to(target)
            except OSError as exc:
                self.skipTest(f"symlink creation is unavailable: {exc}")

            scan = scan_workspace(root)

        self.assertFalse(scan.files)
        self.assertEqual(1, scan.skip_counts["symlink_or_junction"])

    def test_brief_contains_local_evidence(self):
        evidence = self.retriever.search("99.9% uptime")
        _, _, verdict, _, _ = run_court("Did the vendor contractually commit to 99.9% uptime?", evidence, allow_fallback=True, demo_mode=True)
        brief = markdown_brief(verdict, evidence)
        self.assertIn("Cited Evidence", brief)
        self.assertIn("Evidence Ledger", brief)

    def test_evidence_ledger_has_stable_hashes(self):
        entry = self.retriever.evidence[0]
        ledger = self.retriever.evidence_ledger([entry])
        self.assertEqual(entry.citation, ledger[0]["citation"])
        self.assertEqual(64, len(ledger[0]["source_sha256"]))
        self.assertEqual(64, len(ledger[0]["evidence_sha256"]))

    def test_sensitive_scan_redacts_private_material(self):
        fake_token = "ghp_" + "A" * 36
        with TemporaryDirectory() as directory:
            path = Path(directory) / "notes.txt"
            path.write_text(f"temporary token = {fake_token}\n", encoding="utf-8")
            findings = scan_sensitive_paths([path])
        self.assertEqual(1, len(findings))
        self.assertEqual("github token", findings[0].kind)
        self.assertNotIn(fake_token, findings[0].redacted_preview)
        self.assertTrue(findings[0].fingerprint.startswith("sha256:"))

    def test_chinese_credential_request_routes_to_redacted_scanner(self):
        self.assertEqual("sensitive_record_scan", route_workspace_request("帮我找 USA 服务器宝塔账号密码"))

    def test_chinese_assigned_credentials_are_detected_and_redacted(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "usa-server-baota.txt"
            path.write_text("USA 服务器\n宝塔面板 1\n宝塔账号：admin\n密码：Abc12345\n", encoding="utf-8")
            findings = scan_sensitive_paths([path], "帮我找 USA 服务器域名的宝塔面板 1 账号和密码")
            redacted = redact_sensitive_text(path.read_text(encoding="utf-8"))
        self.assertEqual(2, len(findings))
        self.assertTrue(all(item.kind == "assigned secret" for item in findings))
        self.assertTrue(all(item.score > 0 for item in findings))
        self.assertNotIn("admin", redacted)
        self.assertNotIn("Abc12345", redacted)
        self.assertEqual(2, redacted.count("[REDACTED]"))

    def test_sensitive_scan_query_excludes_unrelated_credentials(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            wanted = root / "notes.txt"
            unrelated = root / "database.txt"
            wanted.write_text("USA server\n宝塔面板\n账号：admin\n密码：Abc12345\n", encoding="utf-8")
            unrelated.write_text("database staging\nusername: postgres\npassword: DbSecret9\n", encoding="utf-8")
            findings = scan_sensitive_paths([wanted, unrelated], "找 USA 宝塔的账号密码")
        self.assertEqual({str(wanted.resolve())}, {item.source_path for item in findings})

    def test_redaction_helper_removes_secret_from_display_text(self):
        token = "ghp_" + "B" * 36
        redacted = redact_sensitive_text(f"token={token}")
        self.assertNotIn(token, redacted)
        self.assertIn("[REDACTED", redacted)

    def test_public_embedding_endpoint_is_blocked(self):
        client = LocalEmbeddingClient("vLLM ROCm", "https://api.example.com/v1", "remote")
        with self.assertRaises(ValueError):
            client.embed(["private text"])

    def test_private_network_endpoint_requires_explicit_trust(self):
        client = LocalEmbeddingClient("vLLM ROCm", "http://192.168.1.20:8000/v1", "remote")
        with patch.dict("os.environ", {"CLAIMCOURT_TRUSTED_ENDPOINTS": ""}, clear=False):
            with self.assertRaises(ValueError):
                client.embed(["private text"])

    def test_sensitive_scan_keeps_full_paths_for_duplicate_filenames(self):
        token = "ghp_" + "C" * 36
        with TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "one" / "notes.txt"
            second = root / "two" / "notes.txt"
            first.parent.mkdir()
            second.parent.mkdir()
            first.write_text(token, encoding="utf-8")
            second.write_text(token, encoding="utf-8")
            findings = scan_sensitive_paths([first, second])
        self.assertEqual(2, len(findings))
        self.assertNotEqual(findings[0].source_path, findings[1].source_path)

    def test_default_court_path_rejects_missing_live_judge(self):
        with self.assertRaises(RuntimeError):
            run_court("What was approved?", self.retriever.search("approved"))

    def test_invalid_live_role_schema_is_not_rendered_as_a_verdict(self):
        class BrokenLLM:
            label = "broken-local-model"

            def complete_json(self, system, prompt):
                return [], {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 1, "tokens_per_second": 1.0}

        with self.assertRaises(RuntimeError):
            run_court("What was approved?", self.retriever.search("approved"), BrokenLLM(), allow_fallback=False)

    def test_live_role_name_must_match_requested_role(self):
        class WrongRoleLLM:
            label = "wrong-role-local-model"

            def complete_json(self, system, prompt):
                return {
                    "role": "judge",
                    "position": "Position",
                    "citations": [],
                    "observations": [],
                }, {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 1, "tokens_per_second": 1.0}

        with self.assertRaises(RuntimeError):
            run_court("What was approved?", self.retriever.search("approved"), WrongRoleLLM())

    def test_invalid_live_verdict_enum_is_rejected(self):
        class InvalidVerdictLLM:
            label = "invalid-verdict-local-model"

            def __init__(self):
                self.calls = 0

            def complete_json(self, system, prompt):
                self.calls += 1
                stats = {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 1, "tokens_per_second": 1.0}
                if self.calls <= 2:
                    role = "prosecution" if self.calls == 1 else "defense"
                    return {"role": role, "position": "Position", "citations": [], "observations": []}, stats
                return {
                    "claim": "What was approved?",
                    "verdict": "probably_supported",
                    "confidence": 0.5,
                    "reasoning": "Reasoning",
                    "evidence_citations": [],
                    "contradictions": [],
                    "timeline_events": [],
                    "missing_evidence": [],
                    "recommended_next_action": "Review",
                }, stats

        with self.assertRaises(RuntimeError):
            run_court("What was approved?", self.retriever.search("approved"), InvalidVerdictLLM())

    def test_invalid_model_citations_are_recorded_as_quality_failures(self):
        evidence = self.retriever.search("approved")
        valid = evidence[0].citation

        class CitationLLM:
            label = "citation-local-model"

            def __init__(self):
                self.calls = 0

            def complete_json(self, system, prompt):
                self.calls += 1
                stats = {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 1, "tokens_per_second": 1.0}
                if self.calls <= 2:
                    role = "prosecution" if self.calls == 1 else "defense"
                    return {"role": role, "position": "Position", "citations": [valid, "FAKE-999"], "observations": []}, stats
                return {
                    "claim": "What was approved?",
                    "verdict": "insufficient_evidence",
                    "confidence": 0.5,
                    "reasoning": "Reasoning",
                    "evidence_citations": [valid, "FAKE-999"],
                    "contradictions": [],
                    "timeline_events": [],
                    "missing_evidence": [],
                    "recommended_next_action": "Review",
                }, stats

        prosecution, defense, verdict, telemetry, _ = run_court(
            "What was approved?", evidence, CitationLLM()
        )
        self.assertEqual([valid], prosecution["citations"])
        self.assertEqual([valid], defense["citations"])
        self.assertEqual([valid], verdict["evidence_citations"])
        self.assertTrue(telemetry["model_quality_failures"])
        self.assertTrue(any("FAKE-999" in failure for failure in telemetry["model_quality_failures"]))

    def test_timeline_event_requires_date_event_and_citation(self):
        class InvalidTimelineLLM:
            label = "invalid-timeline-local-model"

            def __init__(self):
                self.calls = 0

            def complete_json(self, system, prompt):
                self.calls += 1
                stats = {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 1, "tokens_per_second": 1.0}
                if self.calls <= 2:
                    role = "prosecution" if self.calls == 1 else "defense"
                    return {"role": role, "position": "Position", "citations": [], "observations": []}, stats
                return {
                    "claim": "What was approved?",
                    "verdict": "insufficient_evidence",
                    "confidence": 0.5,
                    "reasoning": "Reasoning",
                    "evidence_citations": [],
                    "contradictions": [],
                    "timeline_events": [{"event": "Approval discussed", "citation": "E-001"}],
                    "missing_evidence": [],
                    "recommended_next_action": "Review",
                }, stats

        with self.assertRaises(RuntimeError):
            run_court("What was approved?", self.retriever.search("approved"), InvalidTimelineLLM())

    def test_schema_capable_vllm_path_constrains_all_court_roles(self):
        evidence = self.retriever.search("approved")
        valid = evidence[0].citation

        class SchemaLLM:
            label = "schema-local-model"

            def __init__(self):
                self.contracts = []

            def complete_json(self, system, prompt):
                raise AssertionError("schema-capable runtime must use constrained decoding")

            def complete_json_schema(self, system, prompt, name, schema):
                self.contracts.append((name, schema))
                stats = {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 1, "tokens_per_second": 1.0}
                if name == "claimcourt_prosecution":
                    return {"role": "prosecution", "position": "Position", "citations": [valid], "observations": []}, stats
                if name == "claimcourt_defense":
                    return {"role": "defense", "position": "Position", "citations": [valid], "observations": []}, stats
                return {
                    "claim": "What was approved?",
                    "verdict": "insufficient_evidence",
                    "confidence": 0.8,
                    "reasoning": "Reasoning",
                    "evidence_citations": [valid],
                    "contradictions": [],
                    "timeline_events": [],
                    "missing_evidence": [],
                    "recommended_next_action": "Review",
                }, stats

        llm = SchemaLLM()
        _, _, verdict, _, _ = run_court("What was approved?", evidence, llm)
        self.assertEqual("insufficient_evidence", verdict["verdict"])
        self.assertEqual(
            ["claimcourt_prosecution", "claimcourt_defense", "claimcourt_verdict"],
            [name for name, _ in llm.contracts],
        )
        self.assertEqual("number", llm.contracts[-1][1]["properties"]["confidence"]["type"])

    def test_failed_reranker_is_reported_as_disabled(self):
        class BrokenReranker:
            model = "broken-cross-encoder"

            def score(self, query, texts):
                raise RuntimeError("model weights unavailable")

        retriever = LocalRetriever(reranker=BrokenReranker())
        retriever.index_paths((Path(__file__).parent.parent / "demo_corpus").glob("*"))
        results = retriever.search("find the Q4 presentation")
        self.assertTrue(results)
        self.assertEqual("", retriever.reranker_model)
        self.assertIn("model weights unavailable", retriever.reranker_error)
        self.assertFalse(any("cross-encoder reranked" in reason for reason in results[0].match_reasons))

    def test_accumulated_sensitive_scan_keeps_prior_paths(self):
        token = "ghp_" + "D" * 36
        with TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "first.txt"
            second = root / "second.txt"
            first.write_text(token, encoding="utf-8")
            second.write_text(token, encoding="utf-8")
            retriever = LocalRetriever()
            retriever.index_paths([first, second], accumulate=True)
            retriever.index_paths([first], accumulate=True)
            findings = retriever.scan_sensitive_records()
        self.assertEqual(2, len(findings))

    def test_workspace_discovery_and_request_routing(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            nested = root / "nested"
            nested.mkdir()
            (nested / "proposal.pptx").write_bytes(b"placeholder")
            (root / "notes.md").write_text("local note", encoding="utf-8")
            (root / "ignore.exe").write_bytes(b"not evidence")
            self.assertEqual(["notes.md", "proposal.pptx"], sorted(path.name for path in discover_workspace(root)))
        self.assertEqual("file_locator", route_workspace_request("Find the renewal presentation deck"))
        self.assertEqual("sensitive_record_scan", route_workspace_request("Where is my private key?"))
        self.assertEqual("evidence_court", route_workspace_request("Did the vendor approve the delay?"))
        self.assertEqual("file_locator", route_workspace_request("Where is the migration runbook with credential rotation?"))

    def test_vague_request_becomes_transparent_intent_plan(self):
        plan = infer_intent("I wrote a PPT about customer delay and Q4 delivery risk, find it")
        self.assertEqual("locate_artifact", plan.intent)
        self.assertIn(".pptx", plan.artifact_types)
        self.assertIn("delay", plan.topics)
        self.assertIn("risk", plan.topics)
        self.assertIn("Q4", plan.time_hints)
        self.assertIn("slipped", plan.expanded_terms)

    def test_chinese_coursework_request_gets_report_and_topic_metadata(self):
        plan = infer_intent("找出过去一年做过的计算机组成原理课设报告")
        self.assertEqual("locate_artifact", plan.intent)
        self.assertIn(".doc", plan.artifact_types)
        self.assertIn("computer_organization", plan.topics)
        self.assertIn("coursework", plan.topics)
        self.assertIn("past_year", plan.time_hints)

    def test_operating_system_report_request_gets_topic_and_correct_file(self):
        plan = infer_intent("我之前写过一份实验报告，好像是操作系统的，帮我找出来并告诉我写了什么")
        self.assertEqual("locate_artifact", plan.intent)
        self.assertIn("operating_system", plan.topics)
        with TemporaryDirectory() as directory:
            root = Path(directory)
            os_report = root / "课程实验一.docx"
            database_report = root / "课程实验二.docx"
            from docx import Document

            first = Document()
            first.add_heading("操作系统实验报告", level=1)
            first.add_paragraph("实现进程调度、虚拟内存页面置换和文件系统分析。")
            first.save(os_report)
            second = Document()
            second.add_heading("数据库实验报告", level=1)
            second.add_paragraph("实现关系数据库查询、事务和索引。")
            second.save(database_report)
            retriever = LocalRetriever()
            retriever.index_paths([database_report, os_report])
            matches = retriever.locate_files("我之前写过一份实验报告，好像是操作系统的，帮我找出来")
        self.assertEqual(str(os_report.resolve()), matches[0].source_path)

    def test_specific_operating_system_subtopics_are_not_drowned_by_broad_expansion(self):
        cases = (
            ("找基于线程编程的操作系统实验", "线程编程.md"),
            ("找页面置换算法模拟实验报告", "页面置换.md"),
            ("找移动臂磁盘调度算法实验", "磁盘调度.md"),
            ("找包含文件读写的文件操作实验", "文件操作.md"),
        )
        documents = {
            "线程编程.md": "操作系统实验：基于线程编程，创建线程并等待线程结束。",
            "页面置换.md": "操作系统实验：模拟 FIFO 和 LRU 页面置换算法。",
            "磁盘调度.md": "操作系统实验：实现移动臂磁盘调度和 SSTF 算法。",
            "文件操作.md": "操作系统实验：实现文件创建、文件读写和文件管理。",
        }
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for name, text in documents.items():
                (root / name).write_text(text, encoding="utf-8")
            retriever = LocalRetriever()
            retriever.index_paths(root.glob("*.md"))
            for query, expected in cases:
                with self.subTest(query=query):
                    self.assertEqual(expected, retriever.locate_files(query)[0].source)

    def test_filename_path_and_ordinal_metadata_disambiguate_fuzzy_requests(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            dunhuang = root / "作品" / "敦煌古韵"
            career = root / "作品" / "职业规划"
            dunhuang.mkdir(parents=True)
            career.mkdir(parents=True)
            (dunhuang / "最终展示.md").write_text("传统文化主题展示。", encoding="utf-8")
            (career / "最终展示.md").write_text("大学主题展示。", encoding="utf-8")
            (root / "ARM2.md").write_text("ARM 嵌入式实验文档。", encoding="utf-8")
            (root / "ARM3.md").write_text("ARM 嵌入式实验文档。", encoding="utf-8")
            (root / "数据库实验一.md").write_text("数据库 SQL 查询实验。", encoding="utf-8")
            (root / "数据库实验二.md").write_text("数据库 SQL 查询实验。", encoding="utf-8")
            retriever = LocalRetriever()
            retriever.index_paths(root.rglob("*.md"))

            self.assertEqual("最终展示.md", retriever.locate_files("找敦煌古韵的那份资料")[0].source)
            self.assertIn("敦煌古韵", retriever.locate_files("找敦煌古韵的那份资料")[0].source_path)
            self.assertEqual("ARM3.md", retriever.locate_files("找 ARM 实验的第三份文档")[0].source)
            self.assertEqual("数据库实验二.md", retriever.locate_files("找数据库实验二的正式文档")[0].source)

    def test_copy_families_and_document_roles_are_ranked_as_user_requested(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            files = {
                "LINUX1.doc.md": "Linux 操作系统实验一正式报告。",
                "LINUX1.doc - 副本 (2).md": "Linux 操作系统实验一正式报告。",
                "计算机网络实验报告模板.md": "计算机网络实验报告空白模板。",
                "计算机网络实验正式报告.md": "计算机网络实验结果和分析。",
                "课程设计报告模板-成员版.md": "计算机组成与结构课程设计报告模板。",
                "课程设计报告模板-组长版.md": "计算机组成与结构课程设计报告模板。",
            }
            for name, text in files.items():
                (root / name).write_text(text, encoding="utf-8")
            retriever = LocalRetriever()
            retriever.index_paths(root.glob("*.md"))

            linux = retriever.locate_files("找 LINUX1 实验报告")
            linux_family = [match for match in linux if match.source.startswith("LINUX1")]
            self.assertEqual(1, len(linux_family))
            self.assertEqual(1, len(linux_family[0].duplicate_paths))
            self.assertEqual(
                "计算机网络实验正式报告.md",
                retriever.locate_files("找计算机网络实验报告，不要模板")[0].source,
            )
            self.assertEqual(
                "课程设计报告模板-成员版.md",
                retriever.locate_files("找计算机组成与结构课程设计报告模板，成员版，不要组长版")[0].source,
            )

    def test_generic_or_conflicting_coursework_memory_requires_clarification(self):
        self.assertTrue(infer_intent("帮我找我以前写过的那个报告").clarification_needed)
        self.assertTrue(infer_intent("找我去年做的电脑课设").clarification_needed)
        self.assertTrue(infer_intent("找一下那个实验，不记得是数据库还是微机原理了").clarification_needed)

    def test_fuzzy_locator_verbs_and_sensitive_scan_route_deterministically(self):
        locator_queries = (
            "我做过一个数据库实验，好像是第二次，文件名可能不是数据库开头",
            "我以前设计过一个象棋 AI，文件叫什么来着？",
            "我做过五次操作系统实验，帮我列出每次分别讲什么",
        )
        for query in locator_queries:
            with self.subTest(query=query):
                self.assertEqual("file_locator", route_workspace_request(query))
        self.assertEqual(
            "sensitive_record_scan",
            route_workspace_request("Scan my local notes for credentials, return only a masked fingerprint."),
        )
        self.assertEqual(
            "file_locator",
            route_workspace_request("Where is the migration runbook with credential rotation?"),
        )

    def test_multi_experiment_request_prefers_the_requested_series(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for number in (1, 2, 3, 4, 5, 6):
                (root / f"操作系统实验报告（{number}）.md").write_text(
                    f"操作系统第 {number} 次实验内容。", encoding="utf-8"
                )
            (root / "操作系统课设报告.md").write_text("操作系统课程设计综合报告。", encoding="utf-8")
            retriever = LocalRetriever()
            retriever.index_paths(root.glob("*.md"))
            matches = retriever.locate_files("把我操作系统五次实验报告按实验顺序列出来，并分别说明内容", limit=8)
        top_five = {match.source for match in matches[:5]}
        self.assertEqual({f"操作系统实验报告（{number}）.md" for number in (1, 2, 3, 4, 5)}, top_five)

    def test_negated_artifact_type_and_specific_technology_stay_narrow(self):
        slides = infer_intent("找敦煌古韵的幻灯片，不是课程实验报告")
        self.assertEqual({".pptx", ".pdf"}, set(slides.artifact_types))
        verilog = infer_intent("找单周期 CPU 或 Verilog 的课程设计资料")
        self.assertIn("Verilog", verilog.expanded_terms)
        self.assertNotIn("VHDL", verilog.expanded_terms)
        notes = infer_intent("找我整理的计组重点，不要试卷")
        self.assertIn("study_notes", notes.topics)
        self.assertGreater(
            _direct_filename_phrase_bonus("找我整理的计组重点", "计组重点.docx"),
            _direct_filename_phrase_bonus("找我整理的计组重点", "计算机组成原理笔记.pdf"),
        )

    def test_fuzzy_intent_plan_extracts_relations_and_clarification(self):
        plan = infer_intent("找出客户延期之后我做的那个 Q4 风险 PPT")
        self.assertIn("delivery_delay", plan.topics)
        self.assertIn("causal_after", plan.relations)
        self.assertIn("Q4", plan.time_hints)
        self.assertFalse(plan.clarification_needed)
        vague = infer_intent("帮我找一下那个报告")
        self.assertTrue(vague.clarification_needed)
        self.assertTrue(vague.clarification_question)

    def test_local_intent_compiler_merges_model_semantics_without_documents(self):
        class RouterLLM:
            label = "local-intent-router"

            def __init__(self):
                self.prompt = ""

            def complete_json(self, system, prompt):
                self.prompt = prompt
                return {
                    "intent": "locate_artifact",
                    "artifact_types": ["pptx"],
                    "topics": ["customer_delay", "launch_risk"],
                    "entities": ["Northstar"],
                    "time_hints": ["Q4"],
                    "expanded_terms": ["delayed launch", "customer risk"],
                    "search_scope": ["file_name", "full_text", "slide"],
                    "relations": ["causal_after"],
                    "clarification_needed": False,
                    "clarification_question": "",
                    "confidence": 0.91,
                }, {"latency_seconds": 0.2, "first_token_latency_seconds": 0.1, "completion_tokens": 42, "tokens_per_second": 210.0}

        router = RouterLLM()
        plan, telemetry, mode = compile_intent("那个客户延期后的 Q4 PPT", router)
        self.assertEqual("locate_artifact", plan.intent)
        self.assertIn("customer_delay", plan.topics)
        self.assertIn("delivery_delay", plan.topics)
        self.assertIn("causal_after", plan.relations)
        self.assertEqual("local-intent-router", mode)
        self.assertEqual(42, telemetry["completion_tokens"])
        self.assertNotIn("Evidence packet", router.prompt)
        self.assertIn("那个客户延期后的 Q4 PPT", router.prompt)

    def test_intent_compiler_falls_back_with_visible_failure(self):
        class BrokenRouter:
            label = "broken-intent-router"

            def complete_json(self, system, prompt):
                raise RuntimeError("router unavailable")

        plan, telemetry, mode = compile_intent("找出我的操作系统实验报告", BrokenRouter())
        self.assertEqual("locate_artifact", plan.intent)
        self.assertIn("operating_system", plan.topics)
        self.assertEqual("deterministic", mode)
        self.assertIn("router unavailable", telemetry["error"])

    def test_intent_compiler_cannot_suppress_required_clarification(self):
        class OverconfidentRouter:
            label = "overconfident-router"

            def complete_json(self, system, prompt):
                return {
                    "intent": "locate_artifact",
                    "artifact_types": ["all"],
                    "topics": [],
                    "entities": [],
                    "time_hints": [],
                    "expanded_terms": [],
                    "search_scope": ["file_name", "full_text"],
                    "relations": [],
                    "clarification_needed": False,
                    "clarification_question": "",
                    "confidence": 0.95,
                }, {"latency_seconds": 0.1, "completion_tokens": 10, "tokens_per_second": 100.0}

        plan, _, _ = compile_intent("帮我找以前写过的那个报告", OverconfidentRouter())
        self.assertTrue(plan.clarification_needed)
        self.assertTrue(plan.clarification_question)

    def test_intent_compiler_cannot_broaden_sensitive_route(self):
        class OvereagerRouter:
            label = "overeager-router"

            def complete_json(self, system, prompt):
                return {
                    "intent": "sensitive_record_scan",
                    "artifact_types": ["all"],
                    "topics": ["credential"],
                    "entities": [],
                    "time_hints": [],
                    "expanded_terms": [],
                    "search_scope": ["full_text"],
                    "relations": [],
                    "clarification_needed": False,
                    "clarification_question": "",
                    "confidence": 0.99,
                }, {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 10, "tokens_per_second": 100.0}

        plan, _, _ = compile_intent("Where is the migration runbook with credential rotation?", OvereagerRouter())
        self.assertEqual("locate_artifact", plan.intent)

    def test_file_locator_consumes_compiled_expansions(self):
        class RouterLLM:
            label = "local-intent-router"

            def complete_json(self, system, prompt):
                return {
                    "intent": "locate_artifact",
                    "artifact_types": ["pptx"],
                    "topics": ["delivery_risk"],
                    "entities": [],
                    "time_hints": ["Q4"],
                    "expanded_terms": ["customer delay", "launch slipped", "Q4 risk"],
                    "search_scope": ["file_name", "full_text", "slide"],
                    "relations": ["causal_after"],
                    "clarification_needed": False,
                    "clarification_question": "",
                    "confidence": 0.9,
                }, {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 20, "tokens_per_second": 200.0}

        plan, _, _ = compile_intent("就是我说的那个演示稿", RouterLLM())
        matches = self.retriever.locate_files("就是我说的那个演示稿", intent_plan=plan)
        self.assertEqual("Customer_Delivery_Risk_Q4_Final.pptx", matches[0].source)
        self.assertEqual("local-intent-router", self.retriever.last_intent.compiler)

    def test_file_locator_groups_and_explains_semantic_matches(self):
        matches = self.retriever.locate_files("Find the presentation about customer delay and Q4 delivery risk")
        self.assertTrue(matches)
        self.assertTrue(matches[0].evidence)
        self.assertTrue(matches[0].reasons)
        self.assertTrue(any("semantic" in reason for reason in matches[0].reasons))

    def test_signed_addendum_beats_draft_in_same_file_family(self):
        retriever = LocalRetriever()
        corpus = Path(__file__).parent.parent / "championship_corpus"
        retriever.index_paths(discover_workspace(corpus))
        matches = retriever.locate_files("find the signed support addendum with availability exclusions")
        self.assertEqual("Support_Addendum_signed.md", matches[0].source)

    def test_demo_semantic_request_finds_target_presentation(self):
        matches = self.retriever.locate_files("I wrote a PPT about customer delay and Q4 delivery risk, help me find it")
        self.assertEqual("Customer_Delivery_Risk_Q4_Final.pptx", matches[0].source)
        self.assertTrue(any("slide-level" in reason for reason in matches[0].reasons))

    def test_authoritative_version_beats_embedding_bias_within_file_family(self):
        class DraftBiasedEmbedder:
            model = "draft-biased-test"

            def embed(self, texts):
                if len(texts) == 1:
                    return [[1.0, 0.0]]
                return [
                    [1.0, 0.0] if "draft" in text.casefold() else [0.0, 1.0]
                    for text in texts
                ]

        retriever = LocalRetriever(embedder=DraftBiasedEmbedder())
        corpus = Path(__file__).parent.parent / "championship_corpus"
        retriever.index_paths(discover_workspace(corpus))
        query = "I wrote a presentation about the customer's delayed launch and Q4 risk; find it"

        default_match = retriever.locate_files(query, limit=8)[0]
        self.assertEqual("Customer_Delivery_Risk_Q4_Final.pptx", default_match.source)
        self.assertNotIn(default_match.source_path, default_match.duplicate_paths)
        self.assertTrue(any(path.endswith("Customer_Delivery_Risk_Q4_Draft.pptx") for path in default_match.duplicate_paths))
        self.assertIn("authoritative version selected within file family", default_match.reasons)

        draft_match = retriever.locate_files(f"{query} Use the draft version.", limit=8)[0]
        self.assertEqual("Customer_Delivery_Risk_Q4_Draft.pptx", draft_match.source)
        self.assertNotIn(draft_match.source_path, draft_match.duplicate_paths)

    def test_artifact_summary_uses_all_chunks_and_cites_only_selected_file(self):
        match = self.retriever.locate_files("I wrote a PPT about customer delay and Q4 delivery risk, help me find it")[0]
        evidence = [item for item in self.retriever.evidence if item.source_path == match.source_path]

        class SummaryLLM:
            label = "summary-local-model"

            def __init__(self):
                self.prompt = ""

            def complete_json(self, system, prompt):
                self.prompt = prompt
                citations = [item.citation for item in evidence]
                return {
                    "source": match.source,
                    "overview": "A Q4 delivery-risk presentation with customer delay evidence.",
                    "key_points": [{"point": "The deck describes a customer delay and mitigation plan.", "citations": citations[:1]}],
                    "evidence_citations": citations,
                }, {"latency_seconds": 0.1, "first_token_latency_seconds": 0.1, "completion_tokens": 20, "tokens_per_second": 200.0}

        llm = SummaryLLM()
        summary, telemetry = summarize_artifact("What did I write?", match, evidence, llm)
        self.assertEqual(match.source, summary["source"])
        self.assertTrue(summary["key_points"])
        self.assertEqual({item.citation for item in evidence}, set(summary["evidence_citations"]))
        self.assertTrue(all(item.citation in llm.prompt for item in evidence))
        self.assertEqual(20, telemetry["completion_tokens"])

    def test_artifact_summary_requires_live_local_model(self):
        match = self.retriever.locate_files("find the Q4 presentation")[0]
        evidence = [item for item in self.retriever.evidence if item.source_path == match.source_path]
        with self.assertRaises(RuntimeError):
            summarize_artifact("What did I write?", match, evidence, None)

    def test_pptx_index_preserves_slide_locator(self):
        from pptx import Presentation

        with TemporaryDirectory() as directory:
            path = Path(directory) / "renewal-deck.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Renewal Proposal"
            slide.placeholders[1].text = "Customer pricing and Q4 decision timeline"
            presentation.save(path)
            retriever = LocalRetriever()
            retriever.index_paths([path])
        self.assertTrue(retriever.evidence[0].locator.startswith("slide 1"))
        self.assertIn("Renewal Proposal", retriever.evidence[0].text)

    def test_persistent_index_accumulates_and_records_changed_sources(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "decision.md"
            index_file = root / "workspace_index.json"
            source.write_text("first version", encoding="utf-8")
            first = LocalRetriever(index_file)
            first.index_paths([source], accumulate=True)
            first_hash = first.evidence[0].source_sha256
            reopened = LocalRetriever(index_file)
            self.assertEqual(1, len(reopened.evidence))
            source.write_text("changed version", encoding="utf-8")
            reopened.index_paths([source], accumulate=True)
            self.assertNotEqual(first_hash, reopened.evidence[0].source_sha256)
            self.assertTrue(any(item["status"] == "changed" for item in reopened.history))
            self.assertEqual("first version", reopened.archived_versions()[0].text)

    def test_local_embedding_and_query_history_are_optional_but_persistent(self):
        class FakeEmbedder:
            model = "fake-local-embed"

            def embed(self, texts):
                return [[float(len(text)), float(text.lower().count("q4")), float(text.lower().count("delay"))] for text in texts]

        with TemporaryDirectory() as directory:
            root = Path(directory)
            index_file = root / "workspace_index.json"
            retriever = LocalRetriever(index_file, embedder=FakeEmbedder())
            retriever.index_paths((Path(__file__).parent.parent / "demo_corpus").glob("*"))
            matches = retriever.locate_files("find my Q4 delay presentation")
            self.assertIsNotNone(retriever.embedding_matrix)
            self.assertEqual("fake-local-embed", retriever.query_history[-1]["embedding_model"])
            self.assertEqual("locate_artifact", retriever.query_history[-1]["intent"]["intent"])
            reopened = LocalRetriever(index_file)
            self.assertEqual(1, len(reopened.query_history))
            self.assertEqual(matches[0].source, reopened.query_history[0]["results"][0]["source"])

    def test_embedding_client_keeps_inputs_under_model_budget(self):
        client = LocalEmbeddingClient("vLLM ROCm", "http://127.0.0.1:8001/v1", "fake", max_chars=480)
        prepared = client._embedding_text("中" * 900)
        self.assertLessEqual(len(prepared), 480)

    def test_bm25_prioritizes_exact_chinese_phrase(self):
        index = BM25Index(["计算机组成原理课程设计报告", "数据库课程设计报告", "无关的会议记录"])
        scores = index.scores("计算机组成原理课设")
        self.assertGreater(scores[0], scores[1])
        self.assertGreater(scores[0], scores[2])

    def test_failed_local_embedding_endpoint_falls_back_to_hybrid_retrieval(self):
        class FailingEmbedder:
            model = "unavailable-local-embed"

            def embed(self, texts):
                raise ValueError("embedding service unavailable")

        retriever = LocalRetriever(embedder=FailingEmbedder())
        retriever.index_paths((Path(__file__).parent.parent / "demo_corpus").glob("*"))
        matches = retriever.locate_files("find the Q4 delivery presentation")
        self.assertTrue(matches)
        self.assertIsNone(retriever.embedding_matrix)
        self.assertIsNone(retriever.query_history[-1]["embedding_model"])

    def test_fts5_index_is_durable_and_parent_context_is_merged(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "long-record.md"
            source.write_text("intro " * 110 + "TARGET APPROVAL DECISION " + "follow-up " * 110, encoding="utf-8")
            index_file = root / "workspace_index.json"
            retriever = LocalRetriever(index_file)
            retriever.index_paths([source])
            self.assertTrue(retriever.fts5_active)
            self.assertTrue(retriever.evidence[0].parent_id)
            self.assertGreater(retriever.evidence[0].chunk_count, 1)
            reopened = LocalRetriever(index_file)
            results = reopened.search("TARGET APPROVAL DECISION")
            self.assertTrue(results)
            self.assertTrue(reopened.fts5_active)
            self.assertTrue(any(item.context_text for item in results))
            self.assertTrue(any("parent context merged" in reason for reason in results[0].match_reasons))

    def test_local_cross_encoder_reranker_can_promote_a_candidate(self):
        class FakeReranker:
            model = "fake-cross-encoder"

            def score(self, query, texts):
                return [1.0 if "preferred" in text else 0.0 for text in texts]

        with TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "ordinary.md"
            second = root / "preferred.md"
            first.write_text("common workspace record", encoding="utf-8")
            second.write_text("common workspace record preferred answer", encoding="utf-8")
            retriever = LocalRetriever(reranker=FakeReranker())
            retriever.index_paths([first, second])
            result = retriever.search("common workspace record", limit=1)
            self.assertEqual("preferred.md", result[0].source)
            self.assertEqual("fake-cross-encoder", retriever.query_history[-1]["reranker_model"])

    def test_file_decision_abstains_on_weak_or_tied_candidates(self):
        plan = infer_intent("找 2099 年写的量子纠错实验报告")
        weak = FileMatch("weak.md", "/tmp/weak.md", "weak", 0.12, 0.2, (), (), ())
        decision = assess_file_matches("找 2099 年写的量子纠错实验报告", plan, [weak])
        self.assertEqual("no_match", decision.status)

        first = FileMatch("first.md", "/tmp/first.md", "first", 0.8, 0.82, ("concepts: report",), (), ())
        second = FileMatch("second.md", "/tmp/second.md", "second", 0.79, 0.81, ("concepts: report",), (), ())
        decision = assess_file_matches("找操作系统报告", infer_intent("找操作系统报告"), [first, second])
        self.assertEqual("ambiguous", decision.status)
        self.assertTrue(decision.clarification_question)

    def test_file_decision_does_not_collapse_a_requested_series(self):
        matches = [
            FileMatch("one.docx", "/tmp/one.docx", "one", 0.8, 0.82, ("concepts: 操作系统",), (), ()),
            FileMatch("two.docx", "/tmp/two.docx", "two", 0.7, 0.76, ("concepts: 操作系统",), (), ()),
        ]
        query = "把我操作系统五次实验报告按实验顺序列出来，并分别说每份内容"
        self.assertEqual("multiple_matches", assess_file_matches(query, infer_intent(query), matches).status)

    def test_authored_report_inventory_does_not_require_a_topic(self):
        query = "我写过哪些实验报告"
        plan = infer_intent(query)
        self.assertTrue(is_file_inventory_request(query))
        self.assertEqual("locate_artifact", plan.intent)
        self.assertFalse(plan.clarification_needed)

    def test_report_inventory_keeps_only_the_dominant_filled_author(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            paths = []
            for index, topic in enumerate(("thread", "page", "file")):
                path = root / f"owner-{topic}.md"
                path.write_text(f"实验报告 实验学生姓名 李明 学号 100{index} 实验内容 {index}", encoding="utf-8")
                paths.append(path)
            other = root / "other.md"
            other.write_text("实验报告 实验学生姓名 王芳 学号 2001 实验内容", encoding="utf-8")
            paths.append(other)
            retriever = LocalRetriever()
            retriever.index_paths(paths)
            query = "我写过哪些实验报告"
            matches = retriever.locate_files(query, limit=20, intent_plan=infer_intent(query))
            self.assertEqual(3, len(matches))
            self.assertTrue(all(match.source.startswith("owner-") for match in matches))
            self.assertTrue(retriever.last_inventory_identity_inferred)

    def test_report_inventory_applies_the_requested_course_topic(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            paths = []
            for index, topic in enumerate(("thread", "page", "file")):
                path = root / f"os-{topic}.md"
                path.write_text(
                    f"实验报告 课程名称 操作系统 实验项目名称 OS实验{index} 实验学生姓名 李明 学号 100{index}",
                    encoding="utf-8",
                )
                paths.append(path)
            database = root / "database.md"
            database.write_text(
                "实验报告 课程名称 数据库原理 实验项目名称 SQL查询 实验学生姓名 李明 学号 1000",
                encoding="utf-8",
            )
            paths.append(database)
            retriever = LocalRetriever()
            retriever.index_paths(paths)
            query = "操作系统的实验报告有哪些"
            matches = retriever.locate_files(query, limit=20, intent_plan=infer_intent(query))
            self.assertEqual(3, len(matches))
            self.assertTrue(all(match.source.startswith("os-") for match in matches))

    def test_existence_query_requires_topic_role_and_report_constraints(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            database_experiment = root / "database-experiment.md"
            database_experiment.write_text("实验报告 课程名称 数据库原理 实验项目名称 SQL查询", encoding="utf-8")
            os_course_design = root / "os-course-design.md"
            os_course_design.write_text("实验报告 课程名称 操作系统 课程设计 大作业", encoding="utf-8")
            retriever = LocalRetriever()
            retriever.index_paths([database_experiment, os_course_design])
            query = "有没有数据库课设的实验报告"
            self.assertTrue(is_file_existence_request(query))
            matches = retriever.locate_files(query, limit=8, intent_plan=infer_intent(query))
            self.assertEqual([], matches)

    def test_file_inventory_recognizes_natural_list_phrasing(self):
        for query in (
            "列一下操作系统实验报告",
            "把操作系统实验都找出来",
            "我都做过啥操作系统实验",
            "show me every operating system lab report",
        ):
            with self.subTest(query=query):
                self.assertTrue(is_file_inventory_request(query))

    def test_file_locator_enforces_topic_role_negation_and_alternatives(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            paths = {
                "os-lab.md": "实验报告 课程名称 操作系统 实验项目名称 线程同步",
                "database-lab.md": "实验报告 课程名称 数据库原理 实验项目名称 SQL 查询",
                "database-course-design.md": "课程设计报告 课程名称 数据库原理 大作业 数据库系统设计",
                "os-course-design.md": "课程设计报告 课程名称 操作系统 大作业 文件系统设计",
            }
            files = []
            for name, content in paths.items():
                path = root / name
                path.write_text(content, encoding="utf-8")
                files.append(path)
            retriever = LocalRetriever()
            retriever.index_paths(files)

            impossible = "找一下数据库课程设计的实验报告"
            self.assertEqual(
                [],
                retriever.locate_files(impossible, limit=8, intent_plan=infer_intent(impossible)),
            )

            negated = "找实验报告，但不要数据库的"
            negated_matches = retriever.locate_files(negated, limit=8, intent_plan=infer_intent(negated))
            self.assertEqual(["os-lab.md"], [match.source for match in negated_matches])

            alternatives = "找数据库或者操作系统的实验报告"
            alternative_matches = retriever.locate_files(
                alternatives,
                limit=8,
                intent_plan=infer_intent(alternatives),
            )
            self.assertEqual(
                {"database-lab.md", "os-lab.md"},
                {match.source for match in alternative_matches},
            )

    def test_file_decision_does_not_borrow_year_from_lower_candidate(self):
        top = FileMatch(
            "os-report-2024.md",
            "/tmp/os-report-2024.md",
            "os-report-2024",
            0.9,
            0.9,
            ("concepts: operating_system",),
            (),
            (),
        )
        lower = FileMatch(
            "database-report-2025.md",
            "/tmp/database-report-2025.md",
            "database-report-2025",
            0.5,
            0.7,
            ("concepts: database",),
            (),
            (),
        )
        query = "找 2025 年的操作系统实验报告"
        decision = assess_file_matches(query, infer_intent(query), [top, lower])
        self.assertEqual("no_match", decision.status)

    def test_retrieval_feedback_is_local_persistent_and_scope_checked(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "report.md"
            source.write_text("operating system file operation experiment", encoding="utf-8")
            index_file = root / "index.json"
            retriever = LocalRetriever(index_file)
            retriever.index_paths([source])
            retriever.locate_files("find the operating system experiment report")
            query_id = retriever.query_history[-1]["query_id"]
            feedback = retriever.record_feedback(query_id, str(source), True)
            self.assertTrue(feedback["relevant"])
            reopened = LocalRetriever(index_file)
            self.assertEqual(query_id, reopened.feedback_history[-1]["query_id"])
            with self.assertRaises(ValueError):
                reopened.record_feedback(query_id, str(root / "outside.md"), True)

    def test_public_private_eval_report_excludes_queries_paths_and_labels(self):
        private_value = "PRIVATE_QUERY_OR_PATH"
        report = sanitized_report({
            "query_manifest_sha256": "a" * 64,
            "corpus_fingerprint": "b" * 64,
            "elapsed_seconds": 1.0,
            "index": {"indexed_sources": 2},
            "metrics": {"recall_at_5": 1.0},
            "category_metrics": {},
            "retrieval_failures": [{"id": "q1", "query": private_value}],
            "no_answer_failures": [],
            "routing_failures": [],
            "clarification_rows": [],
            "rows": [{"compiler": "deterministic", "expected": private_value}],
        })
        serialized = json.dumps(report)
        self.assertNotIn(private_value, serialized)
        self.assertFalse(report["privacy"]["contains_query_text"])

    def test_scanned_pdf_page_uses_local_ocr_hook(self):
        from pypdf import PdfWriter

        with TemporaryDirectory() as directory:
            path = Path(directory) / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=300)
            with path.open("wb") as stream:
                writer.write(stream)
            with patch("core._ocr_pdf_page", return_value="OCR extracted local text"):
                sections = _document_sections(path)
        self.assertEqual("page 1 (OCR)", sections[0][0])
        self.assertIn("OCR extracted", sections[0][1])

    def test_scanned_pdf_ocr_failure_is_visible_in_diagnostics(self):
        from pypdf import PdfWriter

        with TemporaryDirectory() as directory:
            path = Path(directory) / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=300)
            with path.open("wb") as stream:
                writer.write(stream)
            with patch("core._ocr_pdf_page", return_value=""):
                retriever = LocalRetriever()
                with self.assertRaises(ValueError):
                    retriever.index_paths([path])
        self.assertTrue(any(item["status"] == "empty_page" for item in retriever.ingestion_errors))

    def test_championship_mode_rejects_missing_live_judge(self):
        evidence = self.retriever.search("99.9% uptime")
        with self.assertRaises(RuntimeError):
            run_court(
                "Did the vendor contractually commit to 99.9% uptime?",
                evidence,
                llm=None,
                allow_fallback=False,
            )


if __name__ == "__main__":
    unittest.main()
