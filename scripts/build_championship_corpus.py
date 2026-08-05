"""Build the privacy-safe, multi-file ClaimCourt championship corpus."""

from __future__ import annotations

import argparse
from email.message import EmailMessage
from pathlib import Path


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


def write_email(path: Path, subject: str, body: str) -> None:
    message = EmailMessage()
    message["From"] = "alex@summitsoftware.example"
    message["To"] = "procurement@northstar.example"
    message["Subject"] = subject
    message["Date"] = "Tue, 20 Jan 2026 10:00:00 +0000"
    message.set_content(body.strip() + "\n")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(message.as_bytes())


def write_presentation(path: Path, slides: list[tuple[str, str]]) -> None:
    from pptx import Presentation

    presentation = Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def build(root: Path) -> None:
    if root.exists():
        for path in sorted(root.rglob("*"), reverse=True):
            if path.is_file():
                path.unlink()
            elif path.is_dir():
                path.rmdir()
    root.mkdir(parents=True, exist_ok=True)

    write(root / "01_contracts/Master_Service_Agreement_v1.md", """
    # Master Service Agreement - Draft v1
    Effective date: 2026-01-15.
    The Provider will use commercially reasonable efforts to make the hosted service available.
    This agreement does not establish a service level agreement, guaranteed uptime percentage,
    service credit, or remedy for downtime. Any service level commitment must be set out in a
    separately signed Support Addendum or order form.
    """)
    write(root / "01_contracts/Master_Service_Agreement_signed.md", """
    # Master Service Agreement - Signed
    Effective date: 2026-01-15. Signed by both parties on 2026-01-20.
    The Provider will use commercially reasonable efforts to make the hosted service available.
    This agreement does not establish a service level agreement, guaranteed uptime percentage,
    service credit, or remedy for downtime. Any service level commitment must be set out in a
    separately signed Support Addendum or order form. Marketing statements and sales discussions
    are non-binding unless expressly incorporated into a signed order form.
    """)
    write(root / "01_contracts/Support_Addendum_draft.md", """
    # Support Addendum - Draft
    Version 0.3. Discussion draft dated 2026-02-18; not binding until signed.
    The parties are evaluating a 99.9% availability target after the Q4 failover project.
    Measurement exclusions, service credits, and remedies remain open.
    """)
    write(root / "01_contracts/Support_Addendum_signed.md", """
    # Support Addendum - Signed
    Version 1.0. Signed 2026-04-02. This addendum applies only to the Enterprise Plan.
    Availability is measured at 99.5% monthly, excluding scheduled maintenance and regional outages.
    """)

    write_email(root / "02_sales/Sales_Email_Enterprise_Reliability.eml", "Enterprise-grade reliability for Northstar", """
    Northstar can expect enterprise-grade reliability from Summit's platform. Our team is proud of
    the reliability of the service and will work closely with you during onboarding. This message
    describes the expected experience and is not a signed service-level commitment.
    """)
    write_email(root / "02_sales/Customer_Change_Request.eml", "Northstar requests a Q4 launch change", """
    The customer accepted the revised Q4 launch window in principle, pending a signed change order.
    Please do not treat this email as approval of a contractual SLA or downtime remedy.
    """)

    write(root / "03_meetings/Meeting_Notes_2026-02-03.md", """
    # Reliability Planning Notes - 2026-02-03
    The engineering lead said: "We target 99.9% availability after Q4, provided the regional
    failover project is funded and completed." The account manager confirmed this was a planning
    target, not an executed SLA.
    """)
    write(root / "03_meetings/Steering_Approval_Log.md", """
    # Steering Approval Log
    2026-02-10: Priya Shah approved the migration sequence for a Q4 customer launch.
    2026-02-18: The team approved a two-week schedule change, subject to a signed change order.
    2026-03-01: No one approved a 99.9% contractual uptime commitment.
    """)
    write(root / "03_meetings/Weekly_Status_2026-03-08.md", """
    # Weekly status - 2026-03-08
    Regional failover work is progressing. The Q4 target remains conditional on funding and testing.
    """)

    write_presentation(root / "04_presentations/Customer_Delivery_Risk_Q4_Draft.pptx", [
        ("Customer delivery review - draft", "Internal draft\nPrepared 2026-06-08"),
        ("Customer timeline", "A possible two-week launch delay is under discussion."),
        ("Open questions", "Confirm Q4 ownership and customer approval."),
    ])
    write_presentation(root / "04_presentations/Customer_Delivery_Risk_Q4_Final.pptx", [
        ("Customer delivery review", "Northstar customer launch planning\nFinal review: 2026-06-12"),
        ("Customer timeline change", "Customer requested a two-week delivery delay.\nRevised milestone: Q4 rollout."),
        ("Q4 execution risk", "Risk: regional failover work may slip the Q4 target.\nOwner: Delivery Operations."),
        ("Negotiation options", "Confirm the revised date in a signed change order.\nTrack the Q4 risk in the next steering meeting."),
    ])

    write(root / "05_migration/Migration_Notes.md", """
    # Private migration notes - 2026-05-22
    A synthetic credential fixture is recorded below for the migration owner to rotate. ClaimCourt
    must never display or export this value; only its location and fingerprint may be shown.
    password: DEMO_ONLY_NOT_A_REAL_SECRET
    """)
    write(root / "05_migration/Migration_Runbook.md", """
    # Migration runbook
    Rotate staging credentials after the Q4 rehearsal. The owner is Priya Shah. The final runbook
    contains no customer SLA decision and no approval for a production deployment.
    """)

    for index in range(1, 11):
        write(root / f"90_noise/Unrelated_Project_Note_{index:02d}.md", f"""
        # Unrelated project note {index:02d}
        This document covers routine internal planning, meeting logistics, and office operations.
        It contains no customer delivery decision and no uptime commitment.
        """)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path(__file__).parent.parent / "championship_corpus")
    args = parser.parse_args()
    build(args.output)
    print(f"Built championship corpus at {args.output.resolve()}")
